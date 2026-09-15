#!/usr/bin/env python3
"""Uncommissioned synthetic-layout RFQ demo. No network, OCR or substitutions."""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule

SCHEMA=['code','description_en','description_ar','unit','unit_price_aed','price_approved']
HEADER='CODE | DESCRIPTION | QTY | UNIT'
MARKER='RFQ-DEMO-V1'
LIMIT=Decimal('1000000')
REASONS={
 'UNKNOWN_CODE':('Unknown product code','رمز منتج غير معروف'),
 'AMBIGUOUS_CODE':('Duplicate catalog code','رمز مكرر في الكتالوج'),
 'DESCRIPTION_CONFLICT':('Description conflicts with catalog','الوصف لا يطابق الكتالوج'),
 'UNIT_CONFLICT':('Unit conflicts with catalog','الوحدة لا تطابق الكتالوج'),
 'MISSING_PRICE':('Approved price is missing','السعر المعتمد غير موجود'),
 'INVALID_PRICE':('Price is invalid or out of range','السعر غير صالح أو خارج النطاق'),
 'PRICE_NOT_APPROVED':('Price not marked approved','السعر غير معتمد'),
 'INVALID_QUANTITY':('Quantity must be positive and in range','الكمية يجب أن تكون موجبة وضمن النطاق'),
 'MISSING_CATALOG_DETAIL':('Catalog description or unit missing','وصف الكتالوج أو الوحدة غير موجود'),
}
class InputError(ValueError): pass

def decimal_value(raw,positive=False):
    # Explicit decimal strings only; no binary-float arithmetic, exponent or locale guessing.
    if not re.fullmatch(r'[0-9]+(?:\.[0-9]{1,6})?',str(raw)):
        raise ValueError('Use unsigned decimal digits with at most six decimal places')
    value=Decimal(raw)
    if value>LIMIT or (positive and value<=0): raise ValueError('Out of allowed range')
    return value

def money(quantity,price):
    return (quantity*price).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)

def description_key(text):
    return ' '.join(text.split()).casefold()

def parse_pdf(path):
    if path.stat().st_size>10_000_000: raise InputError('PDF exceeds demo 10 MB limit')
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages)!=1: raise InputError('Agreed demo layout requires exactly one page')
        text=pdf.pages[0].extract_text(x_tolerance=2,y_tolerance=3) or ''
    lines=[line.strip() for line in text.splitlines() if line.strip()]
    if len(lines)<7 or lines[0]!=MARKER or lines[1]!='SYNTHETIC SAMPLE - NOT A CUSTOMER REQUEST':
        raise InputError('Unsupported PDF: expected text-based RFQ-DEMO-V1 synthetic layout')
    if not re.fullmatch(r'RFQ-ID: [A-Za-z0-9_-]{1,40}',lines[2]): raise InputError('Invalid RFQ-ID line')
    if lines[3]!='CURRENCY: AED' or lines[4]!=HEADER or lines[-1]!='END-RFQ':
        raise InputError('Unsupported layout, currency, table header or missing END-RFQ')
    rows=[]
    for line_number,line in enumerate(lines[5:-1],start=6):
        fields=[part.strip() for part in line.split('|')]
        if len(fields)!=4 or not fields[0] or not fields[1] or not fields[3]:
            raise InputError(f'Malformed row at extracted line {line_number}; no row silently skipped')
        if len(fields[0])>40 or len(fields[1])>100 or len(fields[3])>16:
            raise InputError(f'Field exceeds documented layout limits at line {line_number}')
        rows.append(dict(line=len(rows)+1,code=fields[0],description=fields[1],quantity_raw=fields[2],unit=fields[3]))
    if not 1<=len(rows)<=25: raise InputError('Demo layout requires 1 to 25 rows')
    return lines[2].split(': ',1)[1],rows

def read_catalog(path):
    if path.stat().st_size>5_000_000: raise InputError('Catalog exceeds demo 5 MB limit')
    catalog={}
    with path.open(encoding='utf-8-sig',newline='') as handle:
        reader=csv.DictReader(handle)
        if reader.fieldnames!=SCHEMA: raise InputError('Catalog must use exact documented CSV header order')
        for row_number,row in enumerate(reader,start=2):
            if None in row or any(v is None for v in row.values()): raise InputError(f'Malformed catalog row {row_number}')
            row={key:value.strip() for key,value in row.items()}
            if not row['code']: raise InputError(f'Blank catalog code at row {row_number}')
            row['source_row']=row_number
            catalog.setdefault(row['code'],[]).append(row)
    if not catalog: raise InputError('Catalog is empty')
    return catalog

def match_rows(rows,catalog):
    result=[]
    for requested in rows:
        item={**requested,'status':'review','reasons':[],'description_ar':None,'catalog_description':None,
              'catalog_unit':None,'quantity':None,'unit_price_aed':None,'amount_aed':None,'catalog_row':None}
        try: quantity=decimal_value(requested['quantity_raw'],positive=True); item['quantity']=str(quantity)
        except ValueError: quantity=None; item['reasons'].append('INVALID_QUANTITY')
        matches=catalog.get(requested['code'],[])
        if not matches: item['reasons'].append('UNKNOWN_CODE')
        elif len(matches)>1: item['reasons'].append('AMBIGUOUS_CODE')
        else:
            entry=matches[0]
            item.update(description_ar=entry['description_ar'],catalog_description=entry['description_en'],
                        catalog_unit=entry['unit'],catalog_row=entry['source_row'])
            if not all(entry[key] for key in ['description_en','description_ar','unit']):
                item['reasons'].append('MISSING_CATALOG_DETAIL')
            if description_key(requested['description'])!=description_key(entry['description_en']):
                item['reasons'].append('DESCRIPTION_CONFLICT')
            if requested['unit']!=entry['unit']: item['reasons'].append('UNIT_CONFLICT')
            price=None
            if not entry['unit_price_aed']: item['reasons'].append('MISSING_PRICE')
            else:
                try: price=decimal_value(entry['unit_price_aed'])
                except ValueError: item['reasons'].append('INVALID_PRICE')
            if entry['price_approved']!='yes': item['reasons'].append('PRICE_NOT_APPROVED')
            if not item['reasons']:
                item.update(status='matched',unit_price_aed=str(price),amount_aed=str(money(quantity,price)))
        # Any exception holds both price and amount; no suggested substitution is inserted.
        result.append(item)
    return result

def literal(cell,value):
    cell.value=value
    if isinstance(value,str): cell.data_type='s'  # CSV/PDF text can never create an Excel formula.

def export_xlsx(output,rfq_id,items,pdf_path,catalog_path,subtotal):
    wb=Workbook();ws=wb.active;ws.title='Quotation مسودة'
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.sheet_properties.tabColor='294B63'
    ws.freeze_panes='C9';ws.auto_filter.ref=f'A8:I{8+len(items)}'
    rows={
      2:['Quotation draft / مسودة عرض أسعار'],
      3:['DRAFT - NOT FOR CLIENT DELIVERY / مسودة - غير مخصصة للإرسال للعميل'],
      4:['RFQ / الطلب',None,rfq_id,None,'Currency / العملة',None,'AED'],
      5:['Matched / مطابق',None,sum(i['status']=='matched' for i in items),None,'Review / مراجعة',None,sum(i['status']=='review' for i in items)],
      6:['Priced subtotal only / مجموع البنود المسعرة',None,None,None,None,None,None,Decimal(subtotal)],
      7:['Review rows are excluded. No tax, freight or discounts assumed. / بنود المراجعة مستبعدة. دون افتراض ضريبة أو شحن أو خصم.'],
      8:['Line\nالبند','Code\nالرمز','Requested description\nالوصف المطلوب','Catalog Arabic description\nوصف الكتالوج بالعربية','Quantity\nالكمية','Unit\nالوحدة','Unit price (AED)\nسعر الوحدة (درهم)','Line amount (AED)\nقيمة البند (درهم)','Status and exception\nالحالة والاستثناء']
    }
    for row,values in rows.items():
        for col,value in enumerate(values,start=1): literal(ws.cell(row,col),value)
    for row,item in enumerate(items,start=9):
        status='MATCHED DRAFT / مطابق مبدئياً' if not item['reasons'] else '\n'.join(f'{code}: {REASONS[code][0]} / {REASONS[code][1]}' for code in item['reasons'])
        values=[item['line'],item['code'],item['description'],item['description_ar'],
                Decimal(item['quantity']) if item['quantity'] is not None else item['quantity_raw'],item['unit'],
                Decimal(item['unit_price_aed']) if item['unit_price_aed'] is not None else None,
                Decimal(item['amount_aed']) if item['amount_aed'] is not None else None,status]
        for col,value in enumerate(values,start=1): literal(ws.cell(row,col),value)
        ws.row_dimensions[row].height=max(48,34*max(1,len(item['reasons'])))
    widths=[7,18,32,30,13,12,19,19,58]
    for col,width in enumerate(widths,start=1): ws.column_dimensions[chr(64+col)].width=width
    for row in ws:
        for cell in row:
            cell.font=Font(name='Arial',size=11,color='243748')
            cell.alignment=Alignment(vertical='center',wrap_text=cell.row>=8,horizontal='left')
    for row in range(9,9+len(items)):
        for col in [1,5,7,8]: ws.cell(row,col).alignment=Alignment(horizontal='right',vertical='center')
        ws.cell(row,2).alignment=Alignment(horizontal='left',vertical='center',indent=1)
        ws.cell(row,6).alignment=Alignment(horizontal='left',vertical='center',indent=1)
        ws.cell(row,9).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True,indent=1)
        ws.cell(row,4).alignment=Alignment(horizontal='right',vertical='center',wrap_text=True,readingOrder=2)
        ws.cell(row,5).number_format='0' if isinstance(ws.cell(row,5).value,Decimal) and ws.cell(row,5).value==ws.cell(row,5).value.to_integral_value() else '0.######';ws.cell(row,7).number_format='#,##0.00####';ws.cell(row,8).number_format='#,##0.00'
    ws['H6'].number_format='"AED "#,##0.00';ws['H6'].alignment=Alignment(horizontal='right')
    ws['A2'].font=Font(name='Arial',size=16,bold=True,color='243748');ws.row_dimensions[2].height=28
    ws['A3'].font=Font(name='Arial',size=11,bold=True,color='8B3827');ws.row_dimensions[3].height=23
    ws['A6'].font=Font(name='Arial',size=11,bold=True);ws['H6'].font=Font(name='Arial',size=12,bold=True)
    ws.row_dimensions[7].height=24;ws['A7'].font=Font(name='Arial',size=10,color='556273')
    ws.row_dimensions[8].height=46
    for cell in ws[8]:
        cell.fill=PatternFill('solid',fgColor='294B63');cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.conditional_formatting.add(f'A9:I{8+len(items)}',FormulaRule(formula=['LEFT($I9,7)<>"MATCHED"'],fill=PatternFill('solid',fgColor='FFF1D7')))
    footer=10+len(items)
    notes=[
      'Synthetic sample inputs only. Human review has not occurred. / بيانات تجريبية فقط. لم تتم مراجعة بشرية.',
      'Snapshot: edit source inputs and rerun the CLI; spreadsheet edits do not recalculate the matching or Decimal results.',
      f'Sources: {pdf_path.name}; {catalog_path.name}. Exact source hashes appear in the CLI JSON summary.',
      'Rounding: Decimal ROUND_HALF_UP per line to 0.01 AED, then sum rounded matched lines. Exceptions remain blank, not zero.'
    ]
    for offset,note in enumerate(notes): literal(ws.cell(footer+offset,1),note);ws.cell(footer+offset,1).font=Font(name='Arial',size=10,color='556273');ws.row_dimensions[footer+offset].height=22
    ws.print_options.horizontalCentered=True;ws.print_title_rows='1:8';ws.print_area=f'A1:I{footer+len(notes)-1}'
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.page_margins.left=.25;ws.page_margins.right=.25;ws.page_margins.top=.3;ws.page_margins.bottom=.3
    output.parent.mkdir(parents=True,exist_ok=True)
    fd,temp=tempfile.mkstemp(suffix='.xlsx',prefix='.draft-',dir=output.parent);os.close(fd)
    try: wb.save(temp);os.replace(temp,output)
    finally:
        if os.path.exists(temp): os.unlink(temp)

def run(pdf,catalog,output):
    pdf=Path(pdf).resolve();catalog=Path(catalog).resolve();output=Path(output).resolve()
    if output.suffix.lower()!='.xlsx' or output in (pdf,catalog): raise InputError('Output must be a separate .xlsx path')
    rfq_id,requested=parse_pdf(pdf);items=match_rows(requested,read_catalog(catalog))
    subtotal=sum((Decimal(i['amount_aed']) for i in items if i['amount_aed'] is not None),Decimal('0.00'))
    export_xlsx(output,rfq_id,items,pdf,catalog,str(subtotal))
    digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    return dict(status='draft',synthetic_demo=True,human_reviewed=False,rfq_id=rfq_id,currency='AED',rows=len(items),
                matched_rows=sum(i['status']=='matched' for i in items),exception_rows=sum(i['status']=='review' for i in items),
                subtotal_aed=f'{subtotal:.2f}',subtotal_scope='rounded matched lines only; review rows excluded; tax/freight/discounts not included',
                output=str(output),output_sha256=digest(output),pdf_sha256=digest(pdf),catalog_sha256=digest(catalog),items=items)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf',required=True);parser.add_argument('--catalog',required=True);parser.add_argument('--output',required=True)
    args=parser.parse_args()
    try: summary=run(args.pdf,args.catalog,args.output)
    except Exception as error:
        print(json.dumps({'status':'error','error':str(error)},ensure_ascii=False));return 2
    print(json.dumps(summary,ensure_ascii=False));return 0
if __name__=='__main__': raise SystemExit(main())
