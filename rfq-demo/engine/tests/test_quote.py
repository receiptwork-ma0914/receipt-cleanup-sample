import copy,json,subprocess,sys,tempfile,unittest
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
BASE=Path(__file__).resolve().parents[1];sys.path.insert(0,str(BASE))
from quote import money,parse_pdf,read_catalog,match_rows,run,InputError
from generate_fixtures import write_fixture

class QuotationTests(unittest.TestCase):
 def setUp(self): self.catalog=read_catalog(BASE/'fixtures/approved-catalog.csv')
 def test_decimal_half_up_boundary(self):
  self.assertEqual(money(Decimal('0.145'),Decimal('1.00')),Decimal('0.15'))
  self.assertEqual(money(Decimal('1.005'),Decimal('1')),Decimal('1.01'))
  self.assertEqual(money(Decimal('0.144'),Decimal('1')),Decimal('0.14'))
 def test_standard_rows_and_every_requested_exception(self):
  _,rows=parse_pdf(BASE/'fixtures/standard.pdf');items=match_rows(rows,self.catalog)
  self.assertEqual([i['reasons'] for i in items],[[],[],['UNKNOWN_CODE'],['MISSING_PRICE'],['AMBIGUOUS_CODE'],['DESCRIPTION_CONFLICT'],['UNIT_CONFLICT'],[]])
  self.assertEqual([i['amount_aed'] for i in items],['0.15','36.75',None,None,None,None,None,'0.00'])
  for item in items:
   if item['reasons']: self.assertIsNone(item['unit_price_aed']);self.assertIsNone(item['amount_aed'])
 def test_unseen_same_layout_fixture(self):
  with tempfile.TemporaryDirectory() as d:
   summary=run(BASE/'fixtures/unseen.pdf',BASE/'fixtures/approved-catalog.csv',Path(d)/'draft.xlsx')
   self.assertEqual((summary['rows'],summary['matched_rows'],summary['exception_rows']),(4,3,1))
   self.assertEqual(summary['subtotal_aed'],'56.89')
   self.assertEqual([i['amount_aed'] for i in summary['items']],['37.50','1.01','18.38',None])
 def test_exact_code_no_substitution_and_conservative_description(self):
  _,rows=parse_pdf(BASE/'fixtures/standard.pdf');row=rows[0]
  self.assertIn('UNKNOWN_CODE',match_rows([{**row,'code':'cbl-001'}],self.catalog)[0]['reasons'])
  accepted=match_rows([{**row,'description':'  copper   cable 2.5 mm '}],self.catalog)[0]
  self.assertEqual(accepted['status'],'matched')
  self.assertIn('DESCRIPTION_CONFLICT',match_rows([{**row,'description':'Copper cable 4 mm'}],self.catalog)[0]['reasons'])
 def test_bad_quantity_price_and_unapproved_are_row_exceptions(self):
  _,rows=parse_pdf(BASE/'fixtures/standard.pdf');row=rows[0]
  for raw in ['0','-1','NaN','Infinity','1e3','1,5','1000001','0.0000001']:
   self.assertIn('INVALID_QUANTITY',match_rows([{**row,'quantity_raw':raw}],self.catalog)[0]['reasons'])
  for raw in ['NaN','-1','1e3']:
   catalog=copy.deepcopy(self.catalog);catalog['CBL-001'][0]['unit_price_aed']=raw
   self.assertIn('INVALID_PRICE',match_rows([row],catalog)[0]['reasons'])
  catalog=copy.deepcopy(self.catalog);catalog['CBL-001'][0]['price_approved']='no'
  self.assertIn('PRICE_NOT_APPROVED',match_rows([row],catalog)[0]['reasons'])
 def test_one_row_layout_is_supported(self):
  with tempfile.TemporaryDirectory() as d:
   pdf=Path(d)/'one.pdf';write_fixture(pdf,'ONE',[('CBL-001','Copper cable 2.5 mm','0.145','M')])
   self.assertEqual(len(parse_pdf(pdf)[1]),1)
 def test_xlsx_numeric_snapshot_bilingual_blank_exceptions_and_no_formulas(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'draft.xlsx';s=run(BASE/'fixtures/standard.pdf',BASE/'fixtures/approved-catalog.csv',out)
   ws=load_workbook(out,data_only=False).active
   self.assertEqual(ws['H9'].value,0.15);self.assertEqual(ws['H6'].value,36.90)
   self.assertIsNone(ws['G11'].value);self.assertIsNone(ws['H11'].value)
   self.assertEqual(ws['H16'].value,0);self.assertTrue(any('\u0600'<=ch<='\u06ff' for ch in ws['D9'].value))
   self.assertFalse(any(c.data_type=='f' for row in ws for c in row));self.assertEqual(ws.freeze_panes,'C9')
 def test_cli_json_exit_zero_even_with_exceptions_and_fatal_preserves_output(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'draft.xlsx';command=[sys.executable,str(BASE/'quote.py'),'--pdf',str(BASE/'fixtures/standard.pdf'),'--catalog',str(BASE/'fixtures/approved-catalog.csv'),'--output',str(out)]
   p=subprocess.run(command,capture_output=True,text=True);self.assertEqual(p.returncode,0,p.stdout)
   self.assertEqual(json.loads(p.stdout)['exception_rows'],5);before=out.read_bytes()
   bad=Path(d)/'bad.pdf';bad.write_text('not a PDF');command[3]=str(bad)
   p=subprocess.run(command,capture_output=True,text=True);self.assertEqual(p.returncode,2)
   self.assertEqual(json.loads(p.stdout)['status'],'error');self.assertEqual(out.read_bytes(),before)
 def test_formula_like_text_never_becomes_formula(self):
  with tempfile.TemporaryDirectory() as d:
   pdf=Path(d)/'literal.pdf';write_fixture(pdf,'LITERAL',[('=1+1','=HYPERLINK("x")','1','EA')])
   out=Path(d)/'draft.xlsx';run(pdf,BASE/'fixtures/approved-catalog.csv',out)
   ws=load_workbook(out,data_only=False).active
   self.assertEqual(ws['B9'].data_type,'s');self.assertEqual(ws['B9'].value,'=1+1')
 def test_bad_layout_and_catalog_fail_before_export(self):
  with tempfile.TemporaryDirectory() as d:
   bad=Path(d)/'catalog.csv';bad.write_text('code,price\nX,2\n')
   with self.assertRaises(InputError): read_catalog(bad)
   pdf=Path(d)/'badrow.pdf';write_fixture(pdf,'BAD',[('X','Description | extra','1','EA')])
   with self.assertRaises(InputError): parse_pdf(pdf)
if __name__=='__main__': unittest.main(verbosity=2)
